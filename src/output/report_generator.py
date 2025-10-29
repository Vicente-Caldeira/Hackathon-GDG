"""
Generate output reports in JSON and Excel formats.
"""
import json
from datetime import datetime
from pathlib import Path
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from src.validators.consistency_checker import Difference
from config import OUTPUT_DIR


class ReportGenerator:
    """Generate reports from detected differences."""

    @staticmethod
    def generate_json(differences: List[Difference], output_file: Path = None) -> Path:
        """
        Generate JSON report.

        Args:
            differences: List of detected differences
            output_file: Output file path (default: output/differences.json)

        Returns:
            Path to generated file
        """
        if output_file is None:
            output_file = OUTPUT_DIR / "differences.json"

        report = {
            "metadata": {
                "documents": ["test_sample_en", "test_sample_de", "test_sample_lv"],
                "total_differences": len(differences),
                "timestamp": datetime.now().isoformat(),
                "version": "1.0"
            },
            "differences": []
        }

        for i, diff in enumerate(differences, 1):
            report["differences"].append({
                "id": i,
                "paragraph_ids": diff.paragraph_ids,
                "error_type": diff.error_type,
                "severity": diff.severity,
                "values": diff.values,
                "description": diff.description,
                "confidence": diff.confidence
            })

        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(report, f, indent=2, ensure_ascii=False)

        print(f"✓ JSON report saved to: {output_file}")
        return output_file

    @staticmethod
    def generate_excel(differences: List[Difference], output_file: Path = None) -> Path:
        """
        Generate Excel report (similar to errors_test_file.xlsx format).

        Args:
            differences: List of detected differences
            output_file: Output file path (default: output/differences.xlsx)

        Returns:
            Path to generated file
        """
        if output_file is None:
            output_file = OUTPUT_DIR / "differences.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Detected Differences"

        # Define styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        critical_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        medium_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        # Headers
        headers = ["ID", "Paragraph", "EN", "DE", "LV", "Error Type", "Severity", "Description", "Confidence"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows
        for i, diff in enumerate(differences, 1):
            para_ids = diff.paragraph_ids
            # Use EN paragraph ID, or first available
            para_label = para_ids.get('en') or para_ids.get('de') or para_ids.get('lv') or "Unknown"

            row_data = [
                i,
                para_label,
                diff.values.get('en', ''),
                diff.values.get('de', ''),
                diff.values.get('lv', ''),
                diff.error_type,
                diff.severity,
                diff.description,
                f"{diff.confidence:.2f}"
            ]

            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=i+1, column=col, value=value)

                # Apply severity coloring
                if diff.severity == "CRITICAL" and col >= 3 and col <= 5:
                    cell.fill = critical_fill
                elif diff.severity == "MEDIUM" and col >= 3 and col <= 5:
                    cell.fill = medium_fill

        # Adjust column widths
        column_widths = [5, 12, 30, 30, 30, 20, 12, 50, 10]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

        # Freeze header row
        ws.freeze_panes = "A2"

        wb.save(output_file)
        print(f"✓ Excel report saved to: {output_file}")
        return output_file

    @staticmethod
    def generate_summary(differences: List[Difference]) -> str:
        """
        Generate a text summary of findings.

        Args:
            differences: List of detected differences

        Returns:
            Summary string
        """
        total = len(differences)
        by_severity = {}
        by_type = {}

        for diff in differences:
            by_severity[diff.severity] = by_severity.get(diff.severity, 0) + 1
            by_type[diff.error_type] = by_type.get(diff.error_type, 0) + 1

        summary = f"""
Document Consistency Check Summary
{'=' * 60}

Total Differences Found: {total}

By Severity:
"""
        for severity in sorted(by_severity.keys(), key=lambda x: ['CRITICAL', 'MEDIUM', 'LOW'].index(x) if x in ['CRITICAL', 'MEDIUM', 'LOW'] else 999):
            count = by_severity[severity]
            pct = (count / total * 100) if total > 0 else 0
            summary += f"  {severity:12} {count:4} ({pct:.1f}%)\n"

        summary += "\nBy Error Type:\n"
        for error_type in sorted(by_type.keys(), key=lambda x: by_type[x], reverse=True):
            count = by_type[error_type]
            summary += f"  {error_type:25} {count:4}\n"

        return summary


if __name__ == "__main__":
    # Test report generation with dummy data
    from src.validators.consistency_checker import Difference

    dummy_diffs = [
        Difference(
            paragraph_ids={'en': 11, 'de': 11, 'lv': 11},
            error_type="DATE_VALUE",
            severity="MEDIUM",
            values={'en': '16/06/2023', 'de': '15/06/2023', 'lv': '16/06/2023'},
            description="Date differs in German version",
            confidence=0.99
        ),
        Difference(
            paragraph_ids={'en': 13, 'de': 13, 'lv': 13},
            error_type="CURRENCY_MISMATCH",
            severity="CRITICAL",
            values={'en': 'EUR 1500 million', 'de': '1,5 Mrd USD', 'lv': '1500 miljonu EUR'},
            description="Currency mismatch: EUR vs USD",
            confidence=1.0
        )
    ]

    generator = ReportGenerator()
    generator.generate_json(dummy_diffs)
    generator.generate_excel(dummy_diffs)
    print(generator.generate_summary(dummy_diffs))
