#!/usr/bin/env python3
"""
Create a readable report matching the expected format.
Groups all errors by paragraph and shows actual values side-by-side.
"""
import sys
sys.path.insert(0, '.')

from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from src.loaders.json_loader import JSONLoader
from src.alignment.simple_aligner import SimpleAligner
from src.validators.consistency_checker import ConsistencyChecker
from config import INPUT_FILES, OUTPUT_DIR


def create_readable_report():
    """Generate a readable Excel report with actual values."""

    print("Loading documents...")
    documents = JSONLoader.load_all(INPUT_FILES)

    print("Aligning paragraphs...")
    aligner = SimpleAligner()
    aligned = aligner.align_documents(documents)

    print("Checking consistency...")
    checker = ConsistencyChecker()
    differences = checker.check_all(aligned)

    print(f"Found {len(differences)} differences\n")

    # Group by paragraph
    by_paragraph = defaultdict(list)
    for diff in differences:
        para_id = diff.paragraph_ids.get('en') or diff.paragraph_ids.get('de') or diff.paragraph_ids.get('lv')
        if para_id:
            by_paragraph[para_id].append(diff)

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inconsistencies"

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

    # Headers
    headers = ["Paragraph", "EN", "DE", "LV"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 2

    # Process each paragraph
    for para_id in sorted(by_paragraph.keys()):
        diffs = by_paragraph[para_id]

        # Collect all values for this paragraph
        en_values = []
        de_values = []
        lv_values = []

        for diff in diffs:
            vals = diff.values
            en_val = vals.get('en', '')
            de_val = vals.get('de', '')
            lv_val = vals.get('lv', '')

            # Skip generic PRESENT/MISSING rows
            if all(v in ['PRESENT', 'MISSING', ''] for v in [en_val, de_val, lv_val]):
                continue

            # Clean up MISSING markers
            if en_val not in ['MISSING', ''] and en_val not in en_values:
                en_values.append(en_val)
            if de_val not in ['MISSING', ''] and de_val not in de_values:
                de_values.append(de_val)
            if lv_val not in ['MISSING', ''] and lv_val not in lv_values:
                lv_values.append(lv_val)

        # If we have any actual values, write them
        if en_values or de_values or lv_values:
            # Determine max rows needed
            max_rows = max(len(en_values), len(de_values), len(lv_values))

            for i in range(max_rows):
                en_val = en_values[i] if i < len(en_values) else ""
                de_val = de_values[i] if i < len(de_values) else ""
                lv_val = lv_values[i] if i < len(lv_values) else ""

                # Only show paragraph number on first row
                para_label = f"({para_id})" if i == 0 else ""

                ws.cell(row=row_num, column=1, value=para_label)
                ws.cell(row=row_num, column=2, value=en_val)
                ws.cell(row=row_num, column=3, value=de_val)
                ws.cell(row=row_num, column=4, value=lv_val)

                row_num += 1

    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 40

    # Freeze header
    ws.freeze_panes = "A2"

    # Save
    output_file = OUTPUT_DIR / "differences_readable.xlsx"
    wb.save(output_file)

    print(f"✓ Readable report saved to: {output_file}")
    print(f"  Total rows: {row_num - 1}")

    return output_file


if __name__ == "__main__":
    try:
        create_readable_report()
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
